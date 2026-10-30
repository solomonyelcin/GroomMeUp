"""Builds the small synthetic import fixtures used by the tests (original content, no personal data)."""
import openpyxl, docx
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

# ---- xlsx: header on row 2, a hidden sheet, a junk dashboard, a header-less sheet ----
wb = openpyxl.Workbook()
d = wb.active; d.title = "Dashboard"; d["A1"] = "Progress"; d["A2"] = "Total"; d["B2"] = 12
s = wb.create_sheet("Docker Basics")
s["A1"] = "Docker Basics"
s.append(["No", "Category", "Question", "Answer", "Status"])
s.append([1, "Images", "What is a Docker image?", "A read-only template made of layers that contains the app and its dependencies. Containers are started from it.", "Confident"])
s.append([2, "Images", "What is a multi-stage build?", "A Dockerfile with several FROM stages so build tools stay out of the final image, which makes it smaller and safer.", "Pending"])
s.append([3, "Images", "What is a multi-stage build?", "Duplicate of the previous question and must be ignored.", "Pending"])
s.append([4, "Scenario Questions", "Your image is 1.5 GB and slow to deploy. What do you do?", "Use a slim base image, a multi-stage build, a .dockerignore file and remove build caches so only runtime files remain.", "Pending"])
s.append([5, "Images", "Question without an answer", None, "Pending"])
h = wb.create_sheet("Secret notes"); h.append(["Question", "Answer"]); h.append(["Hidden question here?", "Hidden answer text that should never be imported."]); h.sheet_state = "hidden"
g = wb.create_sheet("Git")
g.append(["Q", "A"])
g.append(["What does git rebase do?", "It replays your commits on top of another base commit, producing a linear history instead of a merge commit."])
g.append(["Difference between merge and rebase?", "Merge keeps history and adds a merge commit, rebase rewrites commits to get a straight line. Never rebase shared branches."])
g.append(["What is a fast-forward merge?", "When the target branch has no new commits, Git just moves the branch pointer forward and creates no merge commit."])
nh = wb.create_sheet("No header")
for q, a in [("What is idempotency in HTTP?", "An operation is idempotent if repeating it has the same effect as doing it once, like PUT or DELETE."),
             ("Why use a load balancer?", "It spreads traffic across servers, removes single points of failure and allows rolling deployments."),
             ("What is a CDN used for?", "A CDN caches static content near users to cut latency and reduce load on the origin server.")]:
    nh.append(["", q, a])
wb.save("sample.xlsx")

# ---- docx: Q:/A: markers, a heading, and a table ----
D = docx.Document()
D.add_heading("Redis", level=1)
for q, a in [("What is Redis?", "Redis is an in-memory key-value data store used for caching, queues and sessions."),
             ("What is cache eviction?", "When memory is full Redis removes keys by a policy such as LRU or LFU to make room.")]:
    D.add_paragraph("Q: " + q); D.add_paragraph("A: " + a)
D.add_heading("Kafka", level=1)
D.add_paragraph("Q: What is a consumer group?")
D.add_paragraph("A: A set of consumers that share the partitions of a topic so each message is processed once per group.")
D.add_paragraph("Q: Why are partitions important?")
D.add_paragraph("A: They allow parallelism and ordering per key, and they are the unit of scaling.")
D.add_heading("Networking", level=1)
t = D.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Question"; t.rows[0].cells[1].text = "Answer"
for q, a in [("What is DNS?", "The system that translates domain names into IP addresses."),
             ("What does TLS provide?", "Encryption in transit, server authentication and integrity of the data."),
             ("What is a CIDR block?", "A compact way to write an IP range such as 10.0.0.0/16.")]:
    r = t.add_row().cells; r[0].text = q; r[1].text = a
D.save("sample.docx")

# ---- docx 2: no markers, numbered questions with auto-numbered style ----
E = docx.Document()
E.add_heading("Concurrency", level=1)
E.add_paragraph("What is a deadlock?", style="List Number")
E.add_paragraph("Two or more threads wait forever for locks held by each other. Avoid it with lock ordering and timeouts.")
E.add_paragraph("What is the difference between a mutex and a semaphore?", style="List Number")
E.add_paragraph("A mutex allows one owner at a time, a semaphore allows N concurrent holders and is a counter.")
E.save("sample-numbered.docx")

# ---- text + markdown + csv ----
open("sample.txt", "w").write("""CONCURRENCY

1. What is a race condition?
Two threads touch shared data without coordination, so the result depends on timing.
Fix it with locks, atomic operations or immutable data.

2. Explain thread starvation.
A thread never gets CPU or a lock because others keep winning. Fair locks and priorities help.

TESTING

What is a mock?
A fake collaborator that records calls so a test can verify interactions.

Why write integration tests?
They prove that real components work together, which unit tests with mocks cannot show.
""")
open("sample.md", "w").write("""# Testing

## Unit tests

### What is a stub?
A stub returns canned answers so the code under test can run in isolation.

### What is mutation testing?
It changes the code on purpose and checks that the tests fail, which measures test quality.

## Security

### What is CSRF?
An attack where a logged-in browser is tricked into sending a request the user did not intend. Use anti-forgery tokens and SameSite cookies.
""")
open("sample.csv", "w").write('Topic,Question,Answer\nSQL,"What is a primary key?","A column (or set) that uniquely identifies each row and cannot be NULL."\nSQL,"What is normalization?","Organising tables to reduce duplicate data, e.g. splitting a customer table from an orders table."\nSQL,"Name three join types, with commas?","INNER, LEFT and FULL OUTER joins, each keeping different unmatched rows."\n')
open("notes-only.txt", "w").write("Chapter one explains how load balancing works. It spreads requests over many servers. Health checks remove broken servers from the pool.\n\nChapter two explains caching and expiry.\n")

# ---- pdf with headings, wrapped Q/A lines and a repeating footer ----
c = canvas.Canvas("sample.pdf", pagesize=A4); W, H = A4
def page(title, pairs, n):
    y = H - 60
    c.setFont("Helvetica-Bold", 20); c.drawString(50, y, title); y -= 36
    for q, a in pairs:
        c.setFont("Helvetica-Bold", 11); c.drawString(50, y, "Q: " + q); y -= 18
        c.setFont("Helvetica", 11)
        words = ("A: " + a).split(); line = ""
        for w in words:
            if c.stringWidth(line + " " + w, "Helvetica", 11) > W - 110: c.drawString(50, y, line); y -= 15; line = w
            else: line = (line + " " + w).strip()
        c.drawString(50, y, line); y -= 28
    c.setFont("Helvetica", 9); c.drawString(50, 30, "Interview Notes - Confidential"); c.drawString(W - 90, 30, f"Page {n}")
    c.showPage()
page("Kubernetes", [("What is a Pod?", "The smallest deployable unit in Kubernetes. It holds one or more containers that share network and storage."),
                    ("What does a Deployment do?", "It manages ReplicaSets, gives rolling updates and rollbacks, and keeps the desired number of pods running.")], 1)
page("Linux", [("What is a zombie process?", "A process that has finished but whose parent has not yet read its exit status, so it stays in the process table."),
               ("What does chmod 755 mean?", "The owner can read, write and execute, while group and others can read and execute.")], 2)
page("Cloud", [("What is autoscaling?", "Automatically adding or removing instances based on load so you keep performance and control cost.")], 3)
c.save()
print("fixtures written")
