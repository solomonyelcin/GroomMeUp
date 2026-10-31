"""Questions-only fixtures (no answers): pdf with wrapped lines and headings, docx, xlsx, txt."""
import openpyxl, docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

C_SHARP = ["What is the difference between a struct and a class in C#?",
  "Explain how async and await work under the hood and why blocking on .Result can deadlock an ASP.NET application that uses a synchronization context?",
  "What are the SOLID principles?", "Describe dependency injection lifetimes in ASP.NET Core."]
SQL = ["What is the difference between clustered and non-clustered indexes?",
  "How do you find and fix a slow query in SQL Server?", "Explain isolation levels and the problems each one prevents."]
SYS = ["Design a URL shortener that handles 1000 writes per second.", "Design a rate limiter for a public API."]

def wrap(text, n=70):
    words, lines, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > n: lines.append(cur); cur = w
        else: cur = (cur + " " + w).strip()
    return lines + [cur]

c = canvas.Canvas("questions-only.pdf", pagesize=A4); W, H = A4
y = H - 60
def line(t, size=11, font="Helvetica"):
    global y
    c.setFont(font, size); c.drawString(60, y, t); y -= size + 6
    if y < 80: footer(); c.showPage(); y = H - 60
def footer(): c.setFont("Helvetica", 8); c.drawString(60, 30, "Interview Questions - Page"); 
line("DOTNET INTERVIEW QUESTIONS", 16, "Helvetica-Bold"); y -= 6
line("C# Fundamentals", 13, "Helvetica-Bold")
for i, q in enumerate(C_SHARP, 1):
    ls = wrap(f"{i}. {q}")
    for k, l in enumerate(ls): line(l if k == 0 else "    " + l)
line("SQL Server", 13, "Helvetica-Bold")
for i, q in enumerate(SQL, 1):
    for k, l in enumerate(wrap(f"{i}. {q}")): line(l if k == 0 else "    " + l)
footer(); c.showPage(); y = H - 60
line("System Design", 13, "Helvetica-Bold")
for i, q in enumerate(SYS, 1):
    for k, l in enumerate(wrap(f"{i}. {q}")): line(l if k == 0 else "    " + l)
footer(); c.save()

d = docx.Document()
d.add_heading("Backend questions", 1)
for q in C_SHARP[:3]: d.add_paragraph(q, style="List Number")
d.add_heading("Cloud", 1)
for q in ["What is the difference between a Docker image and a container?", "Explain how Kubernetes handles rolling updates."]: d.add_paragraph(q, style="List Bullet")
d.save("questions-only.docx")

wb = openpyxl.Workbook(); s = wb.active; s.title = "Angular"
s.append(["S.No", "Question", "Category"])
for i, q in enumerate(["What are Angular signals and when do you use them over RxJS?", "Explain change detection and OnPush.", "How do you lazy load a feature module?", "What is the difference between a Subject and a BehaviorSubject?"], 1):
    s.append([i, q, "Core"])
wb.save("questions-only.xlsx")

open("questions-only.txt", "w").write("MICROSERVICES\n\n1. What is the saga pattern?\n2. Explain the transactional outbox pattern and why it matters?\n3. How do you make a consumer idempotent?\n\nREST\n- What is the difference between PUT and PATCH?\n- How would you version a public API?\n")
